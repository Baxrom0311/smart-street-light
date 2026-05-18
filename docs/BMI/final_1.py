#!/usr/bin/env python3
"""
BMI DOCX - TO'LIQ QAYTA YOZISH
- Word Heading styles (dynamic TOC uchun)
- 80 bet asosiy matn (Kirish + I Bob + II Bob + Xulosa)
- 20 bet ilovalar
- 14pt, 1.5 interval, before/after=0
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

# === GLOBAL STYLES ===
style = doc.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(14)
pf = style.paragraph_format
pf.line_spacing = 1.5; pf.space_before = Pt(0); pf.space_after = Pt(0)
pf.first_line_indent = Cm(1.25)

# Heading 1 style (BOB)
h1s = doc.styles['Heading 1']
h1s.font.name = 'Times New Roman'; h1s.font.size = Pt(14); h1s.font.bold = True
h1s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
h1s.paragraph_format.space_before = Pt(0); h1s.paragraph_format.space_after = Pt(0)
h1s.paragraph_format.first_line_indent = Cm(0)
h1s.paragraph_format.line_spacing = 1.5

# Heading 2 style (1.1, 1.2...)
h2s = doc.styles['Heading 2']
h2s.font.name = 'Times New Roman'; h2s.font.size = Pt(14); h2s.font.bold = True
h2s.paragraph_format.space_before = Pt(0); h2s.paragraph_format.space_after = Pt(0)
h2s.paragraph_format.first_line_indent = Cm(0)
h2s.paragraph_format.line_spacing = 1.5

for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(3); s.right_margin = Cm(1.5)

def bob(text):
    """Heading 1 - BOB sarlavhasi (markazda, TOC da ko'rinadi)"""
    p = doc.add_heading(text, level=1)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.color.rgb = None
def h2(text):
    """Heading 2 - bo'lim sarlavhasi (TOC da ko'rinadi)"""
    p = doc.add_heading(text, level=2)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14); r.font.color.rgb = None
def t(x):
    p = doc.add_paragraph(x)
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def img(name, caption):
    if not os.path.exists(IMG + name): return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    run = p.add_run(); run.add_picture(IMG + name, width=Cm(14))
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pc.paragraph_format.first_line_indent = Cm(0)
    pc.paragraph_format.space_before = Pt(0); pc.paragraph_format.space_after = Pt(0)
    rc = pc.add_run(caption); rc.italic = True; rc.font.size = Pt(12); rc.font.name = 'Times New Roman'
def tbl(headers, rows, caption):
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pc.paragraph_format.first_line_indent = Cm(0)
    pc.paragraph_format.space_before = Pt(0); pc.paragraph_format.space_after = Pt(0)
    rc = pc.add_run(caption); rc.italic = True; rc.font.size = Pt(12); rc.font.name = 'Times New Roman'
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = tb.rows[0].cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.size = Pt(12); r.font.name = 'Times New Roman'
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            cell = tb.rows[ri+1].cells[ci]; cell.text = ''
            p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v)); r.font.size = Pt(12); r.font.name = 'Times New Roman'
def code(x):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(x); r.font.name = 'Courier New'; r.font.size = Pt(10)
def add_toc():
    """Word dynamic TOC field"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar1)
    run2 = p.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' TOC \\o "1-2" \\h \\z \\u '
    run2._r.append(instrText)
    run3 = p.add_run()
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fldChar2)
    run4 = p.add_run("(Mundarijani yangilash uchun shu maydonni bosib F9 tugmasini bosing)")
    run4.font.size = Pt(12); run4.font.name = 'Times New Roman'
    run5 = p.add_run()
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run5._r.append(fldChar3)

# ==================== TITUL ====================
for _ in range(4):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.first_line_indent = Cm(0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("O'ZBEKISTON RESPUBLIKASI\nOLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI"); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("MUHAMMAD AL-XORAZMIY NOMIDAGI\nTOSHKENT AXBOROT TEXNOLOGIYALARI UNIVERSITETI"); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run('"KOMPYUTER INJINIRINGI" FAKULTETI\n"DASTURIY INJINIRING" KAFEDRASI'); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
for _ in range(3):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("BITIRUV MALAKAVIY ISHI"); r.bold = True; r.font.size = Pt(16); r.font.name = 'Times New Roman'
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run('Mavzu: "SMART STREET TIZIMIDA ULTRATOVUSH SENSORI ORQALI\nENERGIYA TEJAMKORLIGINI TA\'MINLOVCHI\nYORITISH TIZIMINI ISHLAB CHIQISH"'); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
for _ in range(4):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
t("Bajardi: _______________________")
t("Ilmiy rahbar: _______________________")
for _ in range(5):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("Toshkent — 2026"); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
doc.add_page_break()

# ==================== MUNDARIJA (dynamic) ====================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
r = p.add_run("MUNDARIJA"); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
add_toc()
doc.add_page_break()

print("Titul + TOC done. Saving base...")
doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_final.docx')
print("Base saved. Run part 2 next.")
