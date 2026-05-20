import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

os.chdir(os.path.dirname(os.path.abspath(__file__)))

doc = Document('BMI_final.docx')

# Find indices for 1.2 and 1.3 headings
idx_12 = None
idx_13 = None
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if idx_12 is None and text.startswith('1.2') and len(text) < 150:
        idx_12 = i
    elif idx_12 is not None and idx_13 is None and text.startswith('1.3') and len(text) < 150:
        idx_13 = i
        break

print(f"1.2 heading at index: {idx_12}")
print(f"1.3 heading at index: {idx_13}")

if idx_12 is None or idx_13 is None:
    print("ERROR: Could not find section boundaries!")
    exit(1)

# Remove paragraphs between 1.2 and 1.3 (keep both headings)
body = doc.element.body
elements_to_remove = []
for i in range(idx_12 + 1, idx_13):
    elements_to_remove.append(doc.paragraphs[i]._element)
for el in elements_to_remove:
    body.remove(el)

# Re-find 1.2 heading after removal
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text.startswith('1.2') and len(text) < 150:
        idx_12 = i
        break

IMG_PATH = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

def make_run(paragraph, text, font_name='Times New Roman', font_size=Pt(14), bold=False, italic=False):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    return run

def fmt_paragraph(p, indent=True):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# Load content from separate file
exec(open('content_12.py').read())

# Insert content after 1.2 heading
last_ref = doc.paragraphs[idx_12]
inserted = []

for item in paragraphs_content:
    if item[0] == "text":
        new_p = doc.add_paragraph()
        last_ref._element.addnext(new_p._element)
        fmt_paragraph(new_p, indent=True)
        make_run(new_p, item[1])
        last_ref = new_p
        inserted.append(new_p)
    elif item[0] == "image":
        img_filename = item[1]
        caption_text = item[2]
        # Image paragraph
        img_p = doc.add_paragraph()
        last_ref._element.addnext(img_p._element)
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(6)
        img_p.paragraph_format.space_after = Pt(0)
        img_p.paragraph_format.line_spacing = 1.5
        img_path = os.path.join(IMG_PATH, img_filename)
        if os.path.exists(img_path):
            run = img_p.add_run()
            run.add_picture(img_path, width=Cm(14))
        else:
            make_run(img_p, f'[Rasm: {img_filename}]')
            print(f"  WARNING: Image not found: {img_path}")
        last_ref = img_p
        inserted.append(img_p)
        # Caption
        cap_p = doc.add_paragraph()
        last_ref._element.addnext(cap_p._element)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cap_p.paragraph_format.space_before = Pt(0)
        cap_p.paragraph_format.space_after = Pt(6)
        cap_p.paragraph_format.line_spacing = 1.5
        run = cap_p.add_run(caption_text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.italic = True
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        last_ref = cap_p
        inserted.append(cap_p)

total_chars = sum(len(p.text) for p in inserted)
print(f"\nTotal characters in new 1.2 section: {total_chars}")

doc.save('BMI_final.docx')
print("Saved BMI_final.docx successfully!")
