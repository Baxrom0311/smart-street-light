#!/usr/bin/env python3
"""Add ~28 pages to BMI_full.docx to reach 100+ pages"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

def hc(t, sz=16):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(sz); r.font.name = 'Times New Roman'
def h2(t):
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def t(x):
    p = doc.add_paragraph(x); p.paragraph_format.first_line_indent = Cm(1.25)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def code(x):
    p = doc.add_paragraph(); r = p.add_run(x); r.font.name = 'Courier New'; r.font.size = Pt(10)
    p.paragraph_format.left_indent = Cm(1)
def tbl(headers, rows, caption=""):
    if caption: t(caption)
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, hd in enumerate(headers): tb.rows[0].cells[i].text = hd
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row): tb.rows[ri+1].cells[ci].text = str(v)
    doc.add_paragraph()

# ============ ILOVA G - sensors.cpp to'liq kodi ============
doc.add_page_break()
hc("G ilova. Sensorlar moduli (sensors.cpp)", 14)
doc.add_paragraph()

sensors_code = '''#include "sensors.h"
#include "config.h"
#include <Arduino.h>

// Global o'zgaruvchilar
float currentDistance = 999.0;
int currentLight = 0;
bool motionDetected = false;
unsigned long lastMotionTime = 0;
bool isDark = false;

// Debounce uchun
static int motionCount = 0;
static const int READINGS = 3;
static const int THRESHOLD = 2;

void setupSensors() {
    pinMode(TRIG_PIN, OUTPUT);
    pinMode(ECHO_PIN, INPUT);
    pinMode(LIGHT_PIN, INPUT);
    digitalWrite(TRIG_PIN, LOW);
    Serial.println("[SENSORS] Initialized");
    Serial.print("[SENSORS] TRIG="); Serial.println(TRIG_PIN);
    Serial.print("[SENSORS] ECHO="); Serial.println(ECHO_PIN);
    Serial.print("[SENSORS] LIGHT="); Serial.println(LIGHT_PIN);
}

float readDistance() {
    // TRIG signal yuborish
    digitalWrite(TRIG_PIN, LOW);
    delayMicroseconds(2);
    digitalWrite(TRIG_PIN, HIGH);
    delayMicroseconds(10);
    digitalWrite(TRIG_PIN, LOW);

    // ECHO kutish (timeout 30ms = ~510cm)
    long duration = pulseIn(ECHO_PIN, HIGH, 30000);

    // Timeout - ob'ekt yo'q
    if (duration == 0) return 999.0;

    // Masofa hisoblash: d = t * v / 2
    // v = 343 m/s = 0.0343 cm/us
    float distance = duration * 0.034 / 2.0;

    // Diapazon tekshirish
    if (distance < 2.0) return 2.0;    // Minimal
    if (distance > 450.0) return 999.0; // Maksimal

    return distance;
}

bool checkMotion() {
    motionCount = 0;
    for (int i = 0; i < READINGS; i++) {
        float d = readDistance();
        if (d < DISTANCE_THRESHOLD) {
            motionCount++;
        }
        delay(50); // O'lchovlar orasida pauza
    }
    // Majority voting: 2/3 yoki 3/3
    return motionCount >= THRESHOLD;
}

void readSensors() {
    // Yorug'lik o'qish (12-bit ADC: 0-4095)
    currentLight = analogRead(LIGHT_PIN);

    // Hysteresis bilan kunduz/tun aniqlash
    if (currentLight < DARK_THRESHOLD) {
        isDark = true;
    } else if (currentLight > LIGHT_THRESHOLD) {
        isDark = false;
    }
    // DARK_THRESHOLD va LIGHT_THRESHOLD orasida
    // holat o'zgarmaydi (hysteresis)

    // Masofa o'qish va debounce
    currentDistance = readDistance();

    // Harakat aniqlash (debounce + hold timer)
    if (checkMotion()) {
        lastMotionTime = millis();
        motionDetected = true;
    } else if (millis() - lastMotionTime > HOLD_TIME) {
        motionDetected = false;
    }
    // HOLD_TIME ichida motionDetected = true qoladi
}'''

for line in sensors_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA D - light_control.cpp ============
hc("D ilova. Yoritish boshqaruv moduli (light_control.cpp)", 14)
doc.add_paragraph()

light_code = '''#include "light_control.h"
#include "config.h"
#include "sensors.h"
#include <Arduino.h>

// Global o'zgaruvchilar
bool relayState = false;
String currentMode = "auto";
bool manualLight = false;
String scheduleOn = "18:00";
String scheduleOff = "06:00";

void setupLightControl() {
    pinMode(RELAY_PIN, OUTPUT);
    pinMode(STATUS_LED, OUTPUT);
    digitalWrite(RELAY_PIN, HIGH); // Relay OFF (active LOW)
    digitalWrite(STATUS_LED, LOW);
    Serial.println("[LIGHT] Relay initialized on GPIO " +
                   String(RELAY_PIN));
}

void setRelay(bool state) {
    relayState = state;
    // Relay active LOW - LOW = ON, HIGH = OFF
    digitalWrite(RELAY_PIN, state ? LOW : HIGH);
    // Status LED
    digitalWrite(STATUS_LED, state ? HIGH : LOW);
}

void autoMode() {
    // Kunduz - doim o'chiq
    if (!isDark) {
        setRelay(false);
        return;
    }
    // Tun + harakat = yoqish
    if (motionDetected) {
        setRelay(true);
    } else {
        setRelay(false);
    }
}

void manualMode() {
    setRelay(manualLight);
}

void scheduleMode() {
    // Hozirgi vaqtni olish (NTP yoki millis asosida)
    // Soddalashtirilgan versiya - faqat soat tekshirish
    // To'liq versiyada NTP time ishlatiladi
    setRelay(manualLight); // Placeholder
}

void updateLightControl() {
    if (currentMode == "auto") {
        autoMode();
    } else if (currentMode == "manual") {
        manualMode();
    } else if (currentMode == "schedule") {
        scheduleMode();
    }
}'''

for line in light_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA E - firebase_handler.cpp ============
hc("E ilova. Firebase sinxronizatsiya moduli (firebase_handler.cpp)", 14)
doc.add_paragraph()

firebase_code = '''#include "firebase_handler.h"
#include "config.h"
#include "sensors.h"
#include "light_control.h"
#include <Firebase_ESP_Client.h>
#include <WiFi.h>

FirebaseData fbdo;
FirebaseData stream;
FirebaseAuth auth;
FirebaseConfig config;

unsigned long lastUpdate = 0;
const unsigned long UPDATE_INTERVAL = 3000; // 3 sekund

void setupFirebase() {
    config.api_key = FIREBASE_API_KEY;
    config.database_url = FIREBASE_DB_URL;
    auth.user.email = FIREBASE_EMAIL;
    auth.user.password = FIREBASE_PASSWORD;

    Firebase.begin(&config, &auth);
    Firebase.reconnectWiFi(true);

    // Stream boshlash - control o'zgarishlarini kuzatish
    if (!Firebase.RTDB.beginStream(&stream,
                                    "/device/control")) {
        Serial.println("[FB] Stream error: " +
                       stream.errorReason());
    }
    Serial.println("[FB] Firebase initialized");
}

void readStream() {
    if (!Firebase.ready()) return;
    if (!Firebase.RTDB.readStream(&stream)) return;
    if (!stream.streamAvailable()) return;

    String path = stream.dataPath();
    Serial.println("[FB] Stream: " + path);

    if (path == "/mode" || path.indexOf("mode") >= 0) {
        if (stream.dataType() == "string") {
            currentMode = stream.stringData();
            Serial.println("[FB] Mode: " + currentMode);
        }
    }
    if (path == "/manual_light" ||
        path.indexOf("manual_light") >= 0) {
        if (stream.dataType() == "boolean") {
            manualLight = stream.boolData();
            Serial.println("[FB] Manual: " +
                           String(manualLight));
        }
    }
}

void updateStatus() {
    if (!Firebase.ready()) return;
    if (WiFi.status() != WL_CONNECTED) return;

    Firebase.RTDB.setFloat(&fbdo,
        "/device/status/distance_cm", currentDistance);
    Firebase.RTDB.setInt(&fbdo,
        "/device/status/ambient_light", currentLight);
    Firebase.RTDB.setBool(&fbdo,
        "/device/status/light_on", relayState);
    Firebase.RTDB.setBool(&fbdo,
        "/device/status/motion_detected", motionDetected);
    Firebase.RTDB.setString(&fbdo,
        "/device/status/mode", currentMode);
    Firebase.RTDB.setInt(&fbdo,
        "/device/status/uptime",
        (int)(millis() / 1000));
    Firebase.RTDB.setInt(&fbdo,
        "/device/status/wifi_rssi", WiFi.RSSI());
}

void firebaseLoop() {
    if (WiFi.status() != WL_CONNECTED) return;

    // Stream o'qish (buyruqlar)
    readStream();

    // Status yangilash (har 3 soniyada)
    if (millis() - lastUpdate > UPDATE_INTERVAL) {
        lastUpdate = millis();
        updateStatus();
    }
}'''

for line in firebase_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA J - statistics.cpp ============
hc("J ilova. Statistika moduli (statistics.cpp)", 14)
doc.add_paragraph()

stats_code = '''#include "statistics.h"
#include "config.h"
#include "sensors.h"
#include "light_control.h"
#include <Firebase_ESP_Client.h>
#include <WiFi.h>

extern FirebaseData fbdo;

// Statistika o'zgaruvchilari
unsigned long onDuration = 0;      // Yonish vaqti (ms)
unsigned long lastOnTime = 0;      // Oxirgi yonish boshlanishi
int motionsCount = 0;              // Harakatlar soni
unsigned long lastStatsUpdate = 0;
const unsigned long STATS_INTERVAL = 60000; // 1 daqiqa
bool wasOn = false;

void setupStatistics() {
    onDuration = 0;
    motionsCount = 0;
    lastStatsUpdate = millis();
    Serial.println("[STATS] Initialized");
}

void updateStatistics() {
    // Yonish vaqtini hisoblash
    if (relayState && !wasOn) {
        // Yoqildi
        lastOnTime = millis();
        wasOn = true;
    } else if (!relayState && wasOn) {
        // O'chirildi
        onDuration += millis() - lastOnTime;
        wasOn = false;
    }

    // Harakat sanash
    static bool lastMotion = false;
    if (motionDetected && !lastMotion) {
        motionsCount++;
    }
    lastMotion = motionDetected;

    // Har 1 daqiqada Firebase ga yozish
    if (millis() - lastStatsUpdate > STATS_INTERVAL) {
        lastStatsUpdate = millis();
        saveStats();
    }
}

void saveStats() {
    if (WiFi.status() != WL_CONNECTED) return;
    if (!Firebase.ready()) return;

    // Kunlik statistika
    unsigned long onMin = onDuration / 60000;
    // 12 soat = 720 daqiqa
    int savedPercent = 0;
    if (onMin < 720) {
        savedPercent = ((720 - onMin) * 100) / 720;
    }

    String path = "/history/today";
    Firebase.RTDB.setInt(&fbdo,
        path + "/motions_count", motionsCount);
    Firebase.RTDB.setInt(&fbdo,
        path + "/on_duration_min", (int)onMin);
    Firebase.RTDB.setInt(&fbdo,
        path + "/energy_saved_percent", savedPercent);

    Serial.printf("[STATS] Motions:%d On:%lumin Saved:%d%%\\n",
                  motionsCount, onMin, savedPercent);
}'''

for line in stats_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA Z - config.h ============
hc("Z ilova. Konfiguratsiya fayli (config.h)", 14)
doc.add_paragraph()

config_code = '''#ifndef CONFIG_H
#define CONFIG_H

// ===== PIN KONFIGURATSIYA =====
#define TRIG_PIN        4    // Ultrasonic TRIG
#define ECHO_PIN        16   // Ultrasonic ECHO
#define LIGHT_PIN       34   // TEMT6000 (ADC1_CH6)
#define RELAY_PIN       26   // Relay IN (Active LOW)
#define STATUS_LED      2    // Onboard LED

// ===== SENSOR SOZLAMALARI =====
#define DISTANCE_THRESHOLD  200  // cm (harakat chegarasi)
#define DARK_THRESHOLD      250  // ADC (qorong'u)
#define LIGHT_THRESHOLD     350  // ADC (yorug')
#define HOLD_TIME           5000 // ms (5 sekund)

// ===== WIFI =====
#define WIFI_SSID       "baxrom0311"
#define WIFI_PASS       "12345678"

// ===== FIREBASE =====
#define FIREBASE_API_KEY    "AIzaSyDTGp-CAQmfuhuc4bGYXEaWoEhXnWMpPrU"
#define FIREBASE_DB_URL     "https://smart-street-light-iot-default-rtdb.firebaseio.com"
#define FIREBASE_EMAIL      "device@smartlight.com"
#define FIREBASE_PASSWORD   "SmartLight2026!"

#endif // CONFIG_H'''

for line in config_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA I - React Dashboard kodi ============
hc("I ilova. Dashboard asosiy komponenti (App.jsx)", 14)
doc.add_paragraph()

react_code = '''import { useState, useEffect } from 'react';
import { ref, onValue, set } from 'firebase/database';
import { signInWithEmailAndPassword, signOut } from 'firebase/auth';
import { db, auth } from './firebase';
import Login from './components/Login';
import Dashboard from './components/Dashboard';
import Statistics from './components/Statistics';
import MotionLog from './components/MotionLog';
import Settings from './components/Settings';
import { getMockData, getMockHistory } from './mockData';
import './App.css';

function App() {
  const [user, setUser] = useState(null);
  const [page, setPage] = useState('dashboard');
  const [data, setData] = useState(null);
  const [isLive, setIsLive] = useState(() => {
    return localStorage.getItem('isLive') !== 'false';
  });

  useEffect(() => {
    const unsub = auth.onAuthStateChanged(setUser);
    return unsub;
  }, []);

  useEffect(() => {
    if (!user || !isLive) return;
    const statusRef = ref(db, 'device/status');
    const unsub = onValue(statusRef, (snap) => {
      if (snap.exists()) setData(snap.val());
    });
    return unsub;
  }, [user, isLive]);

  useEffect(() => {
    if (!isLive) setData(getMockData());
  }, [isLive]);

  const handleLogin = async (email, pass) => {
    await signInWithEmailAndPassword(auth, email, pass);
  };

  const togglePower = async () => {
    const newState = !data?.light_on;
    setData(prev => ({...prev, light_on: newState}));
    if (isLive) {
      await set(ref(db, 'device/control/manual_light'), newState);
    }
  };

  const setMode = async (mode) => {
    setData(prev => ({...prev, mode}));
    if (isLive) {
      await set(ref(db, 'device/control/mode'), mode);
    }
  };

  if (!user) return <Login onLogin={handleLogin} />;

  return (
    <div className="app">
      <nav className="sidebar">
        <h2>Smart Light</h2>
        <button onClick={() => setPage('dashboard')}>
          Dashboard
        </button>
        <button onClick={() => setPage('stats')}>
          Statistika
        </button>
        <button onClick={() => setPage('log')}>
          Log
        </button>
        <button onClick={() => setPage('settings')}>
          Sozlamalar
        </button>
        <div className="mode-toggle">
          <span>{isLive ? "LIVE" : "DEMO"}</span>
          <button onClick={() => {
            const next = !isLive;
            setIsLive(next);
            localStorage.setItem('isLive', next);
          }}>Toggle</button>
        </div>
      </nav>
      <main>
        {page === 'dashboard' && (
          <Dashboard data={data}
            onToggle={togglePower}
            onSetMode={setMode} />
        )}
        {page === 'stats' && <Statistics isLive={isLive}/>}
        {page === 'log' && <MotionLog isLive={isLive}/>}
        {page === 'settings' && <Settings isLive={isLive}/>}
      </main>
    </div>
  );
}
export default App;'''

for line in react_code.split('\n'):
    code(line)

doc.add_page_break()

# ============ ILOVA K - platformio.ini ============
hc("K ilova. PlatformIO konfiguratsiya (platformio.ini)", 14)
doc.add_paragraph()

pio_code = '''[env:esp32dev]
platform = espressif32
board = esp32dev
framework = arduino
monitor_speed = 115200
upload_speed = 921600
upload_port = /dev/cu.usbserial-0001
board_build.partitions = huge_app.csv

lib_deps =
    mobizt/Firebase ESP32 Client@^4.4.17

build_flags =
    -DCORE_DEBUG_LEVEL=0'''

for line in pio_code.split('\n'):
    code(line)

doc.add_paragraph()
doc.add_paragraph()
hc("L ilova. Vite konfiguratsiya (vite.config.js)", 14)
doc.add_paragraph()

vite_code = '''import { defineConfig } from 'vite';
import react from '@vitejs/plugin-react';
import { VitePWA } from 'vite-plugin-pwa';

export default defineConfig({
  plugins: [
    react(),
    VitePWA({
      registerType: 'autoUpdate',
      manifest: {
        name: 'Smart Street Light',
        short_name: 'SmartLight',
        theme_color: '#1a1a2e',
        background_color: '#1a1a2e',
        display: 'standalone',
        icons: [{
          src: '/icon-192.png',
          sizes: '192x192',
          type: 'image/png'
        }]
      }
    })
  ]
});'''

for line in vite_code.split('\n'):
    code(line)

# SAVE
doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
print("ILOVALAR added! (~20 pages of code)")
