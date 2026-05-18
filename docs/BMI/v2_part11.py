#!/usr/bin/env python3
"""Push to 100+ pages"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_v2.docx')
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

def h2(t_):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(t_); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def t(x):
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    r = p.add_run(x); r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def tbl(headers, rows, caption):
    pc = doc.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pc.paragraph_format.first_line_indent = Cm(0); pc.paragraph_format.space_before = Pt(0); pc.paragraph_format.space_after = Pt(0)
    rc = pc.add_run(caption); rc.italic = True; rc.font.size = Pt(12); rc.font.name = 'Times New Roman'
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, h in enumerate(headers):
        cell = tb.rows[0].cells[i]; cell.text = ''
        p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h); r.bold = True; r.font.size = Pt(12); r.font.name = 'Times New Roman'
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row):
            cell = tb.rows[ri+1].cells[ci]; cell.text = ''
            p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v)); r.font.size = Pt(12); r.font.name = 'Times New Roman'

doc.add_page_break()
h2("SH' ilova. Mikrokontrollerlar uchun dasturlash tillari va frameworklar tahlili")
t("Embedded tizimlar uchun dasturlash tillari tanlash muhim qaror hisoblanadi, chunki u loyihaning rivojlanish tezligi, samaradorligi va qo'llab-quvvatlash qulayligiga ta'sir qiladi. ESP32 uchun bir nechta dasturlash tillari va frameworklar mavjud bo'lib, ularning har biri o'ziga xos afzallik va kamchiliklarga ega.")
t("C tili embedded tizimlar uchun eng an'anaviy til bo'lib, u to'g'ridan-to'g'ri hardware bilan ishlash imkonini beradi. C da yozilgan kod eng samarali bo'lib, minimal xotira va protsessor resurslarini ishlatadi. Biroq C da dasturlash murakkab va ko'p vaqt talab qiladi, chunki dasturchi xotira boshqaruvini o'zi amalga oshirishi kerak va yuqori darajali abstraksiyalar yo'q.")
t("C++ tili C ning kengaytmasi bo'lib, ob'ektga yo'naltirilgan dasturlash imkoniyatini beradi. Arduino framework C++ da yozilgan va u ESP32 uchun eng mashhur framework hisoblanadi. C++ da klasslar, meros olish va polimorfizm mavjud bo'lib, bu kodni tashkil qilish va qayta ishlatishni osonlashtiradi. Bizning loyihamizda C++ va Arduino framework ishlatilgan.")
t("MicroPython Python tilining mikrokontrollerlar uchun mo'ljallangan versiyasi bo'lib, u ESP32 da ham ishlaydi. MicroPython da dasturlash oson va tez, chunki Python yuqori darajali til va ko'p narsani avtomatik bajaradi. Biroq MicroPython C++ ga nisbatan 10-100 marta sekin ishlaydi va ko'proq xotira sarflaydi. Real vaqt tizimlari uchun bu jiddiy kamchilik.")
t("ESP-IDF Espressif kompaniyasining rasmiy development framework i bo'lib, u C tilida yozilgan va FreeRTOS operatsion tizimiga asoslangan. ESP-IDF eng past darajali kirish imkonini beradi va eng samarali kod yozish mumkin. Biroq u murakkab va o'rganish uchun ko'p vaqt talab qiladi. Arduino framework ESP-IDF ustida qurilgan abstraksiya qatlami bo'lib, u ishlab chiqish tezligini oshiradi.")
t("Bizning loyihamiz uchun Arduino framework tanlandi, chunki u ishlab chiqish tezligini oshiradi, keng jamiyat qo'llab-quvvatlashiga ega va Firebase ESP32 Client kutubxonasi Arduino framework uchun mo'ljallangan. PlatformIO IDE Arduino framework bilan ishlash uchun professional muhit beradi va VS Code bilan integratsiya qilingan.")

tbl(["Til/Framework", "Tezlik", "Xotira", "Ishlab chiqish", "Jamiyat"],
    [["C (ESP-IDF)", "Eng tez", "Eng kam", "Sekin", "O'rtacha"],
     ["C++ (Arduino)", "Tez", "Kam", "Tez", "Juda katta"],
     ["MicroPython", "Sekin", "Ko'p", "Juda tez", "Katta"],
     ["Rust", "Tez", "Kam", "O'rtacha", "Kichik"],
     ["Lua (NodeMCU)", "Sekin", "Ko'p", "Tez", "O'rtacha"]],
    "SH'.1-jadval. ESP32 uchun dasturlash tillari qiyosiy tahlili")

doc.add_page_break()
h2("TS ilova. PlatformIO va loyiha tuzilishi")
t("PlatformIO professional embedded development platform bo'lib, 2014-yilda yaratilgan va hozirda 1000 dan ortiq development board ni qo'llab-quvvatlaydi. U VS Code, CLion va boshqa IDE lar bilan integratsiya qilingan va kutubxonalarni avtomatik boshqaradi. Arduino IDE dan farqli ravishda, PlatformIO multi-environment build, unit testing va CI/CD pipeline larni qo'llab-quvvatlaydi.")
t("PlatformIO loyiha tuzilishi quyidagicha. platformio.ini fayli loyiha konfiguratsiyasini o'z ichiga oladi va unda board turi, framework, kutubxonalar va build flaglar belgilanadi. src papkasida asosiy dastur kodi joylashgan. include papkasida header fayllar joylashgan. lib papkasida lokal kutubxonalar joylashgan. test papkasida unit testlar joylashgan.")
t("Bizning platformio.ini faylimizda quyidagi sozlamalar belgilangan. platform espressif32 ESP32 platformasini ko'rsatadi. board esp32dev standart ESP32 DevKit boardini ko'rsatadi. framework arduino Arduino framework ishlatilishini ko'rsatadi. monitor_speed 115200 serial monitor tezligini belgilaydi. upload_speed 921600 firmware yuklash tezligini belgilaydi. board_build.partitions huge_app.csv katta dasturlar uchun partition sxemasini belgilaydi.")
t("huge_app.csv partition sxemasi standart sxemadan farqli ravishda dastur uchun ko'proq joy ajratadi. Standart sxemada dastur uchun 1.2 MB ajratilgan, huge_app sxemasida esa 3 MB ajratilgan. Bu Firebase kutubxonasi va WiFi stack uchun yetarli joy beradi. Biroq bu sxemada OTA yangilash uchun joy yo'q, shuning uchun kelajakda OTA qo'shilganda boshqa sxemaga o'tish kerak bo'ladi.")
t("Kutubxonalar lib_deps bo'limida belgilanadi. mobizt/Firebase ESP32 Client@^4.4.17 Firebase bilan ishlash uchun asosiy kutubxona. Bu kutubxona Mobizt tomonidan ishlab chiqilgan va GitHub da 3000 dan ortiq yulduzga ega. U Realtime Database, Firestore, Storage va Authentication xizmatlarini qo'llab-quvvatlaydi. Bizning loyihamizda faqat Realtime Database va Authentication ishlatilgan.")

doc.add_page_break()
h2("CH' ilova. Veb-dasturlash texnologiyalari va zamonaviy frontend")
t("Zamonaviy veb-dasturlash bir nechta muhim texnologiyalarga asoslanadi. HTML sahifa tuzilishini belgilaydi, CSS ko'rinishni boshqaradi va JavaScript interaktivlikni ta'minlaydi. Biroq zamonaviy veb-ilovalar uchun bu uchta texnologiya yetarli emas va qo'shimcha vositalar kerak. Framework lar komponent asosida ishlash imkonini beradi, build tool lar kodni optimallashtiriladi va package manager lar kutubxonalarni boshqaradi.")
t("React framework komponent asosida ishlaydi va har bir UI elementi alohida komponent sifatida yaratiladi. Komponent funksiya yoki klass ko'rinishida bo'lishi mumkin va u JSX sintaksisida yoziladi. JSX JavaScript ichida HTML yozish imkonini beradi va u build vaqtida oddiy JavaScript ga kompilyatsiya qilinadi. React ning virtual DOM texnologiyasi faqat o'zgargan qismlarni yangilaydi va bu samaradorlikni oshiradi.")
t("Vite build tool Evan You tomonidan 2020-yilda yaratilgan va u Webpack ga alternativa sifatida ishlab chiqilgan. Vite ning asosiy afzalligi tezlik bo'lib, u ES modules dan foydalanadi va development serverda bundling qilmaydi. Bu Hot Module Replacement ni juda tez qiladi, ya'ni kod o'zgarganda sahifa darhol yangilanadi. Production build da Vite Rollup dan foydalanadi va optimallashtirilgan bundle yaratadi.")
t("npm Node Package Manager Node.js uchun standart package manager bo'lib, u kutubxonalarni o'rnatish, yangilash va o'chirish imkonini beradi. package.json fayli loyihaning barcha bog'liqliklarini ro'yxatga oladi va npm install buyrug'i ularni avtomatik o'rnatadi. Bizning loyihamizda react, firebase, recharts va vite-plugin-pwa kutubxonalari ishlatilgan.")
t("Firebase SDK veb-ilovalar uchun mo'ljallangan JavaScript kutubxonasi bo'lib, u Firebase xizmatlarini brauzerdan ishlatish imkonini beradi. SDK modular arxitekturaga ega va faqat kerakli modullar import qilinadi. Bu bundle hajmini kamaytiradi va ilova tezligini oshiradi. Bizning loyihamizda firebase/app, firebase/database va firebase/auth modullari ishlatilgan.")
t("CSS da zamonaviy layout texnologiyalari ishlatilgan. Flexbox bir o'lchamli layout uchun ishlatiladi va sidebar va main content ni yonma-yon joylashtirish uchun qo'llanilgan. Grid ikki o'lchamli layout uchun ishlatiladi va sensor kartalarini grid ko'rinishida joylashtirish uchun qo'llanilgan. Media queries turli ekran o'lchamlarida turli stillar qo'llash uchun ishlatiladi va responsive dizaynni ta'minlaydi.")

doc.add_page_break()
h2("NG' ilova. Loyihaning iqtisodiy asoslash hujjati")
t("Loyihaning iqtisodiy samaradorligini to'liq baholash uchun barcha xarajatlar va daromadlar hisobga olinishi kerak. Xarajatlar ikki guruhga bo'linadi: bir martalik xarajatlar va doimiy xarajatlar. Bir martalik xarajatlar komponentlar narxi va o'rnatish xarajatlarini o'z ichiga oladi. Doimiy xarajatlar elektr energiyasi va texnik xizmat ko'rsatish xarajatlarini o'z ichiga oladi.")

tbl(["Komponent", "Narx (so'm)", "Narx ($)", "Soni", "Jami"],
    [["ESP32 DevKit", "25,000", "$5", "1", "25,000"],
     ["RCWL-9610A", "5,000", "$1", "1", "5,000"],
     ["TEMT6000", "3,000", "$0.5", "1", "3,000"],
     ["Relay modul", "5,000", "$1", "1", "5,000"],
     ["Simlar va konnektorlar", "5,000", "$1", "1", "5,000"],
     ["Korpus", "10,000", "$2", "1", "10,000"],
     ["USB quvvat manbai", "15,000", "$3", "1", "15,000"],
     ["JAMI", "68,000", "$13.5", "", "68,000"]],
    "NG'.1-jadval. Loyiha komponentlari narxi")

t("Tizimning jami narxi 68000 so'm yoki taxminan 13.5 AQSh dollarini tashkil etadi. Bu tijorat yechimlariga nisbatan 10-30 marta arzon. Philips CityTouch bitta yoritgich uchun 200-500 dollar, Telensa 100-200 dollar talab qiladi. Bizning tizimimiz esa 13.5 dollar bilan bir xil funksionallikni ta'minlaydi.")
t("Yillik doimiy xarajatlar quyidagicha. Tizimning o'zi sarflaydigan elektr energiyasi 0.027 kWh kuniga yoki 9.9 kWh yiliga. Pul hisobida 9.9 × 500 = 4950 so'm yiliga. Firebase bepul rejimda ishlaydi va qo'shimcha xarajat yo'q. Internet uchun mavjud WiFi tarmoqdan foydalaniladi va qo'shimcha xarajat yo'q. Texnik xizmat ko'rsatish minimal bo'lib, faqat sensor tozalash va firmware yangilash kerak.")
t("Yillik daromad yoki tejamkorlik quyidagicha. Bitta 60W yoritgich uchun yillik energiya tejash 212.4 kWh. Pul hisobida 212.4 × 500 = 106200 so'm. Tizim xarajatlarini ayirsak: 106200 - 4950 = 101250 so'm sof yillik tejamkorlik. Tizim narxi 68000 so'm bo'lgani uchun, o'zini oqlash muddati 68000 / 101250 = 0.67 yil yoki taxminan 8 oy.")
t("Investitsiya qaytimi ROI ko'rsatkichi quyidagicha hisoblanadi. Birinchi yil: tejash 101250 - investitsiya 68000 = 33250 so'm foyda. Ikkinchi yildan boshlab: har yili 101250 so'm sof foyda. 5 yillik foyda: 101250 × 5 - 68000 = 438250 so'm. ROI = (438250 / 68000) × 100 = 644 foiz. Bu juda yuqori ko'rsatkich bo'lib, loyihaning iqtisodiy jihatdan juda samarali ekanligini ko'rsatadi.")

tbl(["Yil", "Tejash (so'm)", "Xarajat (so'm)", "Sof foyda", "Jami foyda"],
    [["0 (o'rnatish)", "0", "68,000", "-68,000", "-68,000"],
     ["1", "106,200", "4,950", "101,250", "33,250"],
     ["2", "106,200", "4,950", "101,250", "134,500"],
     ["3", "106,200", "4,950", "101,250", "235,750"],
     ["4", "106,200", "4,950", "101,250", "337,000"],
     ["5", "106,200", "4,950", "101,250", "438,250"]],
    "NG'.2-jadval. 5 yillik iqtisodiy prognoz")

t("Yuqoridagi jadvaldan ko'rinib turibdiki, tizim birinchi yilning 8-oyida o'zini oqlaydi va keyingi yillarda sof foyda beradi. 5 yil davomida bitta yoritgich uchun jami foyda 438250 so'mni tashkil etadi. Bu raqamlar tizimning iqtisodiy jihatdan juda samarali ekanligini yaqqol ko'rsatadi va uni keng miqyosda joriy etish maqsadga muvofiq ekanligini isbotlaydi.")

doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_v2.docx')
print("Done!")
