#!/usr/bin/env python3
"""Final BMI DOCX with real screenshots and diagrams"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document()
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

# Setup styles
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(14)
style.paragraph_format.line_spacing = 1.5
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    s.left_margin = Cm(3); s.right_margin = Cm(1.5)

def h(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(16); r.font.name = 'Times New Roman'
def h2(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def h3(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.italic = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def t(text):
    p = doc.add_paragraph(text); p.paragraph_format.first_line_indent = Cm(1.25)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def img(name, caption="", w=Cm(14)):
    path = IMG + name
    if os.path.exists(path):
        doc.add_picture(path, width=w)
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.font.size = Pt(12); r.italic = True; r.font.name = 'Times New Roman'
def tbl(headers, rows):
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, h in enumerate(headers): tb.rows[0].cells[i].text = h
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row): tb.rows[ri+1].cells[ci].text = str(v)

# Use the existing scripts content but with real images
exec(open('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/generate_docx_1.py').read().split('doc.save')[0].split('doc = Document()')[1])
