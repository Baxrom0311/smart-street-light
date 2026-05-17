#!/usr/bin/env python3
"""II BOB - 2.1 Tizimni loyihalash"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

def hc(t, sz=16):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t); r.bold = True; r.font.size = Pt(sz); r.font.name = 'Times New Roman'
def h2(t):
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def h3(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.italic = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def t(x):
    p = doc.add_paragraph(x); p.paragraph_format.first_line_indent = Cm(1.25)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def code(x):
    p = doc.add_paragraph(); r = p.add_run(x); r.font.name = 'Courier New'; r.font.size = Pt(10)
    p.paragraph_format.left_indent = Cm(1)
def img(name, caption=""):
    if os.path.exists(IMG + name):
        doc.add_picture(IMG + name, width=Cm(14))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.font.size = Pt(12); r.italic = True; r.font.name = 'Times New Roman'
def tbl(headers, rows, caption=""):
    if caption: t(caption)
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, hd in enumerate(headers): tb.rows[0].cells[i].text = hd
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row): tb.rows[ri+1].cells[ci].text = str(v)
    doc.add_paragraph()

hc("II BOB. AMALIY QISM")

h2("2.1. Tizimni loyihalash: texnik vositalar va komponentlarni tanlash")

h3("2.1.1. Mikrokontroller tanlovi va asoslash")
t("Tizimning asosiy boshqaruv elementi sifatida ESP32 DevKit V1 (38-pin) mikrokontrolleri tanlandi. ESP32 — Espressif Systems (Xitoy) kompaniyasi tomonidan 2016-yilda ishlab chiqilgan ikki yadroli mikrokontroller bo'lib, IoT loyihalari uchun eng mashhur platformalardan biri hisoblanadi [20]. Tanlash jarayonida bir nechta alternativ variantlar ko'rib chiqildi:")

tbl(["Parametr", "Arduino Uno", "ESP8266", "ESP32", "STM32F103", "Raspberry Pi Pico"],
    [["Protsessor", "ATmega328P", "Tensilica L106", "Xtensa LX6 (2x)", "ARM Cortex-M3", "ARM Cortex-M0+ (2x)"],
     ["Chastota", "16 MHz", "80 MHz", "240 MHz", "72 MHz", "133 MHz"],
     ["RAM", "2 KB", "80 KB", "520 KB", "20 KB", "264 KB"],
     ["Flash", "32 KB", "4 MB", "4 MB", "64 KB", "2 MB"],
     ["WiFi", "Yo'q", "802.11 b/g/n", "802.11 b/g/n", "Yo'q", "Yo'q (W variant ha)"],
     ["Bluetooth", "Yo'q", "Yo'q", "BLE 4.2", "Yo'q", "Yo'q"],
     ["ADC", "6ch / 10-bit", "1ch / 10-bit", "18ch / 12-bit", "10ch / 12-bit", "3ch / 12-bit"],
     ["GPIO", "14", "11", "34", "37", "26"],
     ["Narx", "~$5", "~$3", "~$5", "~$3", "~$4"]],
    "2.1-jadval. Mikrokontrollerlarning qiyosiy tahlili")

t("ESP32 tanlash sabablari: 1) Ichki WiFi moduli — tashqi WiFi shield kerak emas, bu narxni va murakkablikni kamaytiradi; 2) Ikki yadroli protsessor (240 MHz) — bitta yadro sensorlarni o'qiydi, ikkinchisi WiFi/Firebase bilan ishlaydi, bu non-blocking arxitekturani ta'minlaydi; 3) 12-bitli ADC (18 kanal) — TEMT6000 yorug'lik sensorini aniq o'qish uchun; 4) 520 KB RAM — Firebase kutubxonasi va WiFi stack uchun yetarli; 5) 3.3V logika — RCWL-9610A bilan to'g'ridan-to'g'ri ulash mumkin; 6) Arzon narx (~$5) — loyiha byudjetiga mos; 7) PlatformIO va Arduino framework qo'llab-quvvatlashi — tez ishlab chiqish.")

h3("2.1.2. Tizim arxitekturasi")
t("Tizim uch qatlamli arxitekturaga ega: Hardware (ESP32 + sensorlar + relay), Cloud (Firebase Realtime Database + Authentication + Hosting), va Client (React PWA Dashboard). Har bir qatlam mustaqil ishlaydi va boshqa qatlamlar bilan aniq belgilangan interfeys orqali aloqa qiladi.")

img("diagram_architecture.png", "2.1-rasm. Tizimning umumiy arxitekturasi")

t("2.1-rasmda tizimning umumiy arxitekturasi tasvirlangan. ESP32 sensorlardan ma'lumot o'qiydi, qaror qabul qiladi va relay orqali yoritgichni boshqaradi. Har 3 soniyada status ma'lumotlarini Firebase ga yuboradi. Dashboard Firebase dan real vaqt rejimida ma'lumot oladi va foydalanuvchi buyruqlarini Firebase ga yozadi. ESP32 stream orqali buyruqlarni qabul qiladi.")

h3("2.1.3. Elektr sxemasi va pin konfiguratsiyasi")
t("ESP32 DevKit 38-pin versiyasida jami 34 ta GPIO pin mavjud, biroq ularning hammasi ham erkin ishlatilishi mumkin emas. Ba'zi pinlar 'strapping pins' deb ataladi va boot jarayonida maxsus vazifa bajaradi (GPIO 0, 2, 5, 12, 15). Ba'zilari esa faqat input sifatida ishlaydi (GPIO 34, 35, 36, 39). Shu sababli, pin tanlashda ehtiyotkorlik zarur [20].")

tbl(["ESP32 Pin", "Komponent", "Signal turi", "Yo'nalish", "Izoh"],
    [["GPIO 4", "RCWL-9610A TRIG", "Digital", "OUTPUT", "Trigger signal"],
     ["GPIO 16", "RCWL-9610A ECHO", "Digital", "INPUT", "Echo signal"],
     ["GPIO 34", "TEMT6000 OUT", "Analog (ADC1_CH6)", "INPUT", "Faqat input pin"],
     ["GPIO 26", "Relay IN", "Digital", "OUTPUT", "Active LOW"],
     ["GPIO 2", "Onboard LED", "Digital", "OUTPUT", "Status indikator"]],
    "2.2-jadval. ESP32 pin tayinlash")

t("Pin tanlash asoslari: GPIO 4 — strapping pin emas, digital output uchun mos; GPIO 16 — UART2 RX, lekin biz UART2 ishlatmaymiz; GPIO 34 — faqat input, ADC1 kanalida (WiFi bilan conflict yo'q); GPIO 26 — DAC2 pin, lekin biz DAC ishlatmaymiz, digital output uchun mos; GPIO 2 — onboard LED, debug uchun qulay.")

h3("2.1.4. Quvvat ta'minoti va energiya iste'moli")
t("Tizimning umumiy energiya iste'moli quyidagicha hisoblanadi:")

tbl(["Komponent", "Ish toki", "Kutish toki", "Izoh"],
    [["ESP32 (WiFi active)", "150 mA", "20 mA", "WiFi uzatish vaqtida"],
     ["RCWL-9610A", "2 mA", "2 mA", "Doimiy"],
     ["TEMT6000", "0.1 mA", "0.1 mA", "Doimiy"],
     ["Relay modul", "70 mA", "0 mA", "Faqat yoqilganda"],
     ["JAMI (max)", "222.1 mA", "22.1 mA", ""]],
    "2.3-jadval. Tizim energiya iste'moli")

t("Tizim 5V / 1A USB quvvat manbai orqali ta'minlanadi. Maksimal iste'mol 222 mA bo'lib, bu 5V da 1.11W ga teng. Kuniga tizimning o'zi sarflaydigan energiya: 1.11W × 24h = 26.6 Wh = 0.027 kWh. Bu yoritgichning o'zi sarflaydigan energiyaga (720 Wh/kun) nisbatan 27 marta kam — ya'ni tizimning o'zi energiya tejash samaradorligiga deyarli ta'sir qilmaydi.")

h3("2.1.5. Dasturiy ta'minot platformasi va kutubxonalar")
t("Firmware ishlab chiqish uchun PlatformIO IDE tanlandi. PlatformIO — bu professional embedded development platform bo'lib, Arduino framework ni qo'llab-quvvatlaydi va VS Code bilan integratsiya qilingan [24]. Arduino IDE dan farqli ravishda, PlatformIO kutubxonalarni avtomatik boshqaradi, multi-environment build ni qo'llab-quvvatlaydi va professional CI/CD pipeline larni sozlash imkonini beradi.")

tbl(["Kutubxona", "Versiya", "Vazifasi"],
    [["Arduino Framework", "ESP32 v2.x", "Asosiy framework"],
     ["Firebase ESP32 Client", "v4.4.17", "Firebase RTDB bilan aloqa"],
     ["WiFi.h", "Built-in", "WiFi ulanish"],
     ["Preferences.h", "Built-in", "NVS (Non-Volatile Storage)"],
     ["DNSServer.h", "Built-in", "Captive portal uchun"],
     ["WebServer.h", "Built-in", "AP rejimda web server"]],
    "2.4-jadval. Ishlatilgan kutubxonalar")

t("Web dashboard uchun React 19 + Vite 8 tanlandi. React — Meta (Facebook) tomonidan ishlab chiqilgan UI kutubxona bo'lib, komponent asosida ishlaydi [22]. Vite — yangi avlod build tool bo'lib, Hot Module Replacement (HMR) va tezkor build ni ta'minlaydi [25]. Firebase SDK 10 real vaqt ma'lumot almashish uchun, Recharts esa grafiklar uchun ishlatildi.")

doc.add_page_break()
doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
print("2.1 done (~10 pages)")
