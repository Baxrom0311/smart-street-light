#!/usr/bin/env python3
"""BMI DOCX Generator - Smart Street Light diplom ishi"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

doc = Document()

# Styles
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(14)
style.paragraph_format.line_spacing = 1.5

# Margins
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

def add_heading_centered(text, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 0 else 14)
    run.font.name = 'Times New Roman'
    return p

def add_normal(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Cm(1.25)
    return p

def add_image(path, width=Cm(14)):
    if os.path.exists(path):
        doc.add_picture(path, width=width)
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER

IMG_PATH = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

# ===== TITUL SAHIFA =====
for _ in range(3):
    doc.add_paragraph()

add_heading_centered("O'ZBEKISTON RESPUBLIKASI OLIY TA'LIM, FAN VA INNOVATSIYALAR VAZIRLIGI")
add_heading_centered("MUHAMMAD AL-XORAZMIY NOMIDAGI\nTOSHKENT AXBOROT TEXNOLOGIYALARI UNIVERSITETI")
doc.add_paragraph()
add_heading_centered('"KOMPYUTER INJINIRINGI" FAKULTETI\n"DASTURIY INJINIRING" KAFEDRASI')

for _ in range(2):
    doc.add_paragraph()

add_heading_centered("BITIRUV MALAKAVIY ISHI")
doc.add_paragraph()
add_heading_centered('Mavzu:\n"SMART STREET TIZIMIDA ULTRATOVUSH SENSORI ORQALI\nENERGIYA TEJAMKORLIGINI TA\'MINLOVCHI\nYORITISH TIZIMINI ISHLAB CHIQISH"')

for _ in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.add_run("Bajardi: _______________________").font.size = Pt(14)
p = doc.add_paragraph()
p.add_run("Ilmiy rahbar: _______________________").font.size = Pt(14)

for _ in range(4):
    doc.add_paragraph()

add_heading_centered("Toshkent — 2026")
doc.add_page_break()

# ===== MUNDARIJA =====
add_heading_centered("MUNDARIJA")
doc.add_paragraph()

mundarija = [
    ("KIRISH", "4"),
    ("I BOB. SMART STREET TIZIMIDA YORITISH TIZIMLARI TIZIMLI TAHLILI VA MASALANING QO'YILISHI", ""),
    ("1.1. Smart Street tizimlarida energiya tejamkor yoritish tizimlari va ularning ahamiyati", "9"),
    ("1.2. Ultratovush sensorlarning ishlash prinsipi va texnik xususiyatlari", "19"),
    ("1.3. Masalaning qo'yilishi", "32"),
    ("II BOB. AMALIY QISM", ""),
    ("2.1. Tizimni loyihalash: texnik vositalar va komponentlarni tanlash", "36"),
    ("2.2. Masofani aniqlash algoritmi: ultrasonic sensor yordamida obyektni aniqlash", "45"),
    ("2.3. Yoritish tizimini boshqarish dasturi va tizimning ishlash jarayoni va natijalar tahlili", "54"),
    ("XULOSA", "65"),
    ("FOYDALANILGAN ADABIYOTLAR RO'YXATI", "68"),
    ("ILOVALAR", "70"),
]

for title, page in mundarija:
    p = doc.add_paragraph()
    if page:
        p.add_run(f"{title} {'.' * (60 - len(title))} {page}")
    else:
        run = p.add_run(title)
        run.bold = True

doc.add_page_break()

# Save intermediate
doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/bmi_draft.docx')
print("Part 1 done: titul + mundarija")
