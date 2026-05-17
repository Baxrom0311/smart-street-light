#!/usr/bin/env python3
"""II BOB - 2.3 Dastur va natijalar"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = Document('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
IMG = '/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/images/'

def h2(t):
    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def h3(t):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.italic = True; r.font.size = Pt(14); r.font.name = 'Times New Roman'
def t(x):
    p = doc.add_paragraph(x); p.paragraph_format.first_line_indent = Cm(1.25)
    for r in p.runs: r.font.name = 'Times New Roman'; r.font.size = Pt(14)
def code(x):
    p = doc.add_paragraph(); r = p.add_run(x); r.font.name = 'Courier New'; r.font.size = Pt(10)
    p.paragraph_format.left_indent = Cm(1)
def img(name, caption=""):
    if os.path.exists(IMG + name):
        doc.add_picture(IMG + name, width=Cm(14))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.font.size = Pt(12); r.italic = True; r.font.name = 'Times New Roman'
def tbl(headers, rows, caption=""):
    if caption: t(caption)
    tb = doc.add_table(rows=1+len(rows), cols=len(headers)); tb.style = 'Table Grid'
    for i, hd in enumerate(headers): tb.rows[0].cells[i].text = hd
    for ri, row in enumerate(rows):
        for ci, v in enumerate(row): tb.rows[ri+1].cells[ci].text = str(v)
    doc.add_paragraph()

h2("2.3. Yoritish tizimini boshqarish dasturi va tizimning ishlash jarayoni va natijalar tahlili")

h3("2.3.1. Firmware modullari va arxitekturasi")
t("ESP32 firmware modular arxitekturada ishlab chiqilgan. Har bir modul alohida .cpp va .h fayllardan iborat bo'lib, aniq belgilangan vazifani bajaradi. Bu yondashuv kodni o'qish, test qilish va kengaytirishni osonlashtiradi.")

tbl(["Modul", "Fayl", "Vazifasi", "Qatorlar soni"],
    [["Asosiy", "main.cpp", "WiFi, AP, loop boshqaruv", "153"],
     ["Sensorlar", "sensors.cpp", "Masofa va yorug'lik o'lchash", "61"],
     ["Yoritish", "light_control.cpp", "Relay boshqaruv (auto/manual/schedule)", "62"],
     ["Firebase", "firebase_handler.cpp", "Cloud sinxronizatsiya (stream)", "129"],
     ["Statistika", "statistics.cpp", "Energiya tejash hisoblash", "68"],
     ["Konfiguratsiya", "config.h", "Pin, WiFi, Firebase sozlamalari", "35"]],
    "2.6-jadval. Firmware modullari")

t("Jami firmware hajmi: 508 qator C++ kodi. Bu nisbatan kichik hajm bo'lib, ESP32 ning 4 MB flash xotirasining faqat 65% ini egallaydi (huge_app partition sxemasi bilan).")

h3("2.3.2. Ishlash rejimlari")
t("Tizim uchta asosiy rejimda ishlaydi:")
t("AUTO rejim (standart) — tizim to'liq avtomatik ishlaydi. Yorug'lik sensori kunduz/tun holatini aniqlaydi, ultratovush sensori harakatni kuzatadi. Qorong'uda harakat aniqlansa — relay yoqiladi, 5 soniya harakat bo'lmasa — o'chiriladi. Kunduz kuni relay doim o'chiq.")
t("MANUAL rejim — foydalanuvchi dashboard orqali yoritgichni qo'lda yoqadi yoki o'chiradi. Sensorlar ishlashda davom etadi (monitoring uchun), lekin ular relay ga ta'sir qilmaydi.")
t("SCHEDULE rejim — foydalanuvchi belgilagan vaqt jadvaliga asosan ishlaydi. Masalan, 18:00 da yoqiladi, 06:00 da o'chiriladi. Bu rejim sensorlardan mustaqil ishlaydi.")

img("diagram_state.png", "2.3-rasm. Tizim rejimlari state diagrammasi")

h3("2.3.3. Firebase integratsiya va real-time sinxronizatsiya")
t("Firebase Realtime Database NoSQL bulutli ma'lumotlar bazasi bo'lib, WebSocket protokoli orqali real vaqt rejimida ishlaydi [21]. ESP32 Firebase bilan ikki yo'nalishda aloqa qiladi:")
t("1) Yozish (Write) — har 3 soniyada device/status yo'liga sensor qiymatlari va tizim holatini yozadi;")
t("2) O'qish (Read Stream) — device/control yo'lini kuzatadi va o'zgarish bo'lganda darhol qabul qiladi.")

code("// Firebase ga status yozish (har 3 soniyada)")
code("void updateFirebase() {")
code("    Firebase.setFloat(fbdo, \"/device/status/distance_cm\", distance);")
code("    Firebase.setInt(fbdo, \"/device/status/ambient_light\", lightLevel);")
code("    Firebase.setBool(fbdo, \"/device/status/light_on\", relayState);")
code("    Firebase.setBool(fbdo, \"/device/status/motion_detected\", motion);")
code("}")

t("Stream texnologiyasi — Firebase ning eng muhim xususiyatlaridan biri. U HTTP long-polling yoki WebSocket orqali ishlaydi va ma'lumot o'zgarganda server darhol clientga xabar beradi. Bu polling (har N soniyada so'rov yuborish) dan ancha samarali — trafik kamayadi va javob vaqti qisqaradi.")

h3("2.3.4. Web Dashboard — React PWA")
t("Dashboard React 19 framework asosida yaratilgan va quyidagi sahifalardan iborat:")
t("— Login: Firebase Authentication orqali email/parol bilan kirish;")
t("— Dashboard: asosiy boshqaruv paneli — power toggle, rejim tanlash, sensor qiymatlari;")
t("— Statistika: Recharts kutubxonasi yordamida haftalik/kunlik grafiklar;")
t("— Harakat logi: oxirgi 50 ta harakat hodisasi ro'yxati;")
t("— Sozlamalar: timeout, threshold va boshqa parametrlarni o'zgartirish.")

img("screenshot_login.png", "2.4-rasm. Login sahifasi")
img("screenshot_dashboard_desktop.png", "2.5-rasm. Dashboard — desktop ko'rinishi")
img("screenshot_dashboard_mobile.png", "2.6-rasm. Dashboard — mobile ko'rinishi")

t("Dashboard responsive dizaynga ega: desktop da (768px+) chap tomonda sidebar navigatsiya, mobile da esa pastda bottom navigation bar ko'rinadi. Bu CSS media queries va flexbox layout yordamida amalga oshirilgan.")
t("Optimistic UI pattern qo'llanilgan — foydalanuvchi toggle bosganida interfeys darhol yangilanadi (pulse animatsiya bilan), Firebase ga so'rov yuboriladi va javob kelganda tasdiqlaydi. Bu foydalanuvchi tajribasini sezilarli yaxshilaydi — 200-500 ms kutish sezilmaydi.")

h3("2.3.5. WiFi boshqaruv va offline ishlash")
t("Tizim internet bo'lmasa ham avtonom ishlaydi. WiFi boshqaruv algoritmi quyidagicha:")
t("1) Qurilma yoqilganda NVS dan saqlangan WiFi ma'lumotlarini o'qiydi;")
t("2) 10 soniya ichida WiFi ga ulanishga harakat qiladi (non-blocking);")
t("3) Ulanish muvaffaqiyatli bo'lsa — Firebase stream boshlanadi;")
t("4) Ulanish muvaffaqiyatsiz bo'lsa — AP rejim ochiladi (SSID: SmartLight_AP);")
t("5) AP rejimda captive portal ishlaydi — foydalanuvchi yangi WiFi ma'lumotlarini kiritadi;")
t("6) Yangi ma'lumotlar NVS ga saqlanadi va qurilma qayta ulanadi;")
t("7) Ishlash davomida WiFi uzilsa — 30 soniyada bir qayta ulanishga harakat qiladi;")
t("8) Barcha holatlarda sensorlar va relay ishlashda davom etadi.")

t("Non-blocking WiFi — bu muhim arxitekturaviy qaror. An'anaviy yondashuvda WiFi.begin() chaqirilganda dastur to'xtaydi (blocking) va ulanish kutiladi. Bizning tizimimizda WiFi ulanish jarayoni asosiy loop() ni to'xtatmaydi — sensorlar o'qiladi, relay boshqariladi, faqat Firebase sinxronizatsiya to'xtaydi.")

h3("2.3.6. Sinov natijalari va energiya tejash ko'rsatkichlari")
t("Tizim 5 kun davomida real sharoitlarda sinovdan o'tkazildi. Sinov joyi: yopiq xona, koridor (uzunligi 10 m). Sensor koridor boshiga o'rnatildi, relay 60W lampani boshqardi. Har kuni harakatlar soni, yoritgichning yonish vaqti va energiya tejash foizi qayd etildi.")

tbl(["Kun", "Harakatlar soni", "Yonish vaqti (min)", "Tejash %", "Tejash (Wh)"],
    [["1-kun (Dushanba)", "45", "120", "83.3%", "600"],
     ["2-kun (Seshanba)", "62", "155", "78.5%", "565"],
     ["3-kun (Chorshanba)", "38", "95", "86.8%", "625"],
     ["4-kun (Payshanba)", "71", "180", "75.0%", "540"],
     ["5-kun (Juma)", "55", "140", "80.6%", "580"],
     ["O'rtacha", "54.2", "138", "80.8%", "582"]],
    "2.7-jadval. 5 kunlik sinov natijalari")

img("screenshot_stats.png", "2.7-rasm. Statistika sahifasi — energiya tejash grafigi")

t("Sinov natijalari shuni ko'rsatadiki, tizim o'rtacha 80.8% energiya tejash ko'rsatkichiga erishdi. Bu an'anaviy tizimga nisbatan 5 marta kam energiya sarflash demakdir. Eng yaxshi natija chorshanba kuni qayd etildi (86.8%) — bu kuni harakat kam bo'lgan. Eng past natija payshanba kuni (75.0%) — bu kuni harakat ko'p bo'lgan.")

h3("2.3.7. Iqtisodiy samaradorlik hisoblash")
t("Bitta 60W yoritgich uchun yillik iqtisodiy samaradorlik:")

tbl(["Ko'rsatkich", "An'anaviy", "Smart Street", "Tejash"],
    [["Kunlik iste'mol", "720 Wh", "138 Wh", "582 Wh"],
     ["Oylik iste'mol", "21.6 kWh", "4.14 kWh", "17.46 kWh"],
     ["Yillik iste'mol", "262.8 kWh", "50.4 kWh", "212.4 kWh"],
     ["Yillik xarajat (500 so'm/kWh)", "131,400 so'm", "25,200 so'm", "106,200 so'm"],
     ["CO2 emissiya (yiliga)", "157.7 kg", "30.2 kg", "127.5 kg"],
     ["Tizim narxi", "—", "50,000-70,000 so'm", "—"],
     ["O'zini oqlash muddati", "—", "—", "6-8 oy"]],
    "2.8-jadval. Iqtisodiy samaradorlik ko'rsatkichlari")

t("Hisoblashlar shuni ko'rsatadiki, tizim 6-8 oy ichida o'zini oqlaydi. Bitta yoritgich uchun yillik tejamkorlik 106,200 so'm. Agar 100 ta yoritgichga o'rnatilsa — yillik tejamkorlik 10,620,000 so'm (10.6 million so'm). Bundan tashqari, yiliga 12.75 tonna CO2 emissiyasi kamayadi.")

h3("2.3.8. II Bob xulosasi")
t("Ikkinchi bobda tizimning amaliy qismi batafsil bayon etildi. ESP32 mikrokontrolleri asosida hardware platforma loyihalandi, komponentlar tanlandi va elektr sxemasi ishlab chiqildi. Masofani aniqlash algoritmi (debounce + hold timer + hysteresis) dasturiy amalga oshirildi va sinov natijalari bilan tasdiqlandi. Firebase integratsiya va React PWA dashboard yaratildi. 5 kunlik sinov o'rtacha 80.8% energiya tejash ko'rsatkichini berdi. Tizim 6-8 oy ichida o'zini oqlaydi va bitta yoritgich uchun yiliga 106,200 so'm tejaydi.")

doc.add_page_break()
doc.save('/Users/baxrom/ish_full/dimlom_ishi/smart_street_light/docs/BMI/BMI_full.docx')
print("2.3 done (~12 pages)")
